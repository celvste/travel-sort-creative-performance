#!/usr/bin/env python3
"""
applovin_week.py - weekly AppLovin creative pull for Travel Sort UA.

Pulls asset-level performance from AppLovin's Asset Reporting API and writes a
finished weekly tab per campaign - same columns, formulas and formatting as the
tabs in the master workbooks - ready to copy straight in.

Usage
-----
    python3 applovin_week.py                          # every campaign, last completed week
    python3 applovin_week.py --campaign BLD
    python3 applovin_week.py --campaign IAP --week-end 2026-09-15
    python3 applovin_week.py --list-campaigns
    python3 applovin_week.py --csv                    # also drop raw htmls/videos/images.csv

Layout
------
    applovin-ua/
      applovin_week.py
      config.json
      totals/          <- drop this week's Total sheet exports here
      weeks/           <- output, one folder per week

Added Date
----------
Read from the Total sheet of each campaign, exported as CSV into totals/.
Export each master workbook's Total tab (File > Download > CSV) and drop the
file in unchanged - the script matches it to a campaign by the label appearing
in the filename, so _Travel_Sort__D28_IAP_..._-_Total.csv lands on IAP.

Refresh those files each week before running. A creative missing from the
Total sheet gets its Added Date looked up from the API instead: the earliest
day it drew an impression in that campaign, shown in a blue cell.

That is a first-SERVE date, not an upload date, and the API only reaches
back 45 days. Anything older stays amber and blank for manual entry.
Pass --no-date-lookup to skip the extra calls.

Config (./config.json)
----------------------
    {
      "report_key": "...",
      "campaigns": {
        "IAP":     {"id": "93707df1...", "highlight": 2000},
        "BLD_IOS": {"id": "81e0c8bb..."},
        "BLD_AND": {"id": "7c31f90a..."},
        "ADS":     {"id": ["50aa720e...", "9b2d41ff..."]}
      }
    }

`id` is a campaign ID, an exact campaign name, or a list of either.

A campaign split by platform (BLD iOS and BLD Android) is two campaigns in
AppLovin, so give each its own label - they have different CPIs and should
not share a Spending % denominator. Use a list of ids only when you want
the platforms genuinely combined in one tab.

Labels are matched against the filenames in totals/, so a label that is a
prefix of another (BLD vs BLD_IOS) needs its Total export named to match. `highlight` is the spend
threshold for the green fill (default 1000). An optional `workbook` path makes
the script copy last week's row order from that master workbook; without it,
rows follow the order in the Total sheet. report_key may also come from
APPLOVIN_REPORT_KEY.

Notes
-----
* The API is UTC; yesterday's data is stable after 06:00 UTC (14:00 SGT).
* assetAnalyticsReport only accepts dates within the last 45 days.
"""

import argparse
import csv
import glob
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://r.applovin.com/assetAnalyticsReport"
PAGE_SIZE = 1000

HEADER = ["Creative Type", "Added Date", "Creatives", "Spend", "CTR", "Impression", "Click", "Spending %"]
WIDTHS = [11, 12, 52, 12, 11, 12, 11, 12]
HEADER_FILL = PatternFill("solid", fgColor="FF812520")
HEADER_FONT = Font(name="Cambria", sz=11, bold=True, color="FFFFFFFF")
BODY_FONT = Font(name="Calibri", sz=11)
AMBER = PatternFill("solid", fgColor="FFFFE699")
BLUE = PatternFill("solid", fgColor="FFD6E4F5")   # date derived from first serve
LOOKBACK_DAYS = 44                                # API allows 45; stay inside
GREEN = "FFD9F5D6"
TAB_RE = re.compile(r"^\d{2}\.\d{2}\.\d{2}-\d{2}\.\d{2}\.\d{2}$")
TYPE_ORDER = ["Video", "Playable", "Image"]


# ------------------------------------------------------------------ config


def load_config(path):
    path = os.path.expanduser(path)
    cfg = json.load(open(path)) if os.path.exists(path) else {}
    key = os.environ.get("APPLOVIN_REPORT_KEY") or cfg.get("report_key")
    if not key:
        sys.exit(f'No report key. Add "report_key" to {path} or set APPLOVIN_REPORT_KEY.')
    cfg["report_key"] = key
    cfg["campaigns"] = {
        label: ({"id": spec} if isinstance(spec, str) else dict(spec))
        for label, spec in (cfg.get("campaigns") or {}).items()
    }
    return cfg


def campaign_ids(spec, label):
    """`id` may be one identifier or a list of them.

    A platform-split campaign (BLD iOS + BLD Android) is two campaigns in
    AppLovin. Give the label both ids and their asset rows are summed, or
    give each platform its own label to keep them apart.
    """
    raw = spec.get("id") or label
    return [str(x).strip() for x in (raw if isinstance(raw, list) else [raw]) if str(x).strip()]


# ------------------------------------------------------------------- dates


def last_tuesday(today=None):
    today = today or date.today()
    offset = (today.weekday() - 1) % 7 or 7
    return today - timedelta(days=offset)


def week_label(start, end):
    return f"{start:%y.%m.%d}-{end:%y.%m.%d}"


# ------------------------------------------------------------------- fetch


def fetch(params, key):
    rows, offset = [], 0
    while True:
        query = dict(params, api_key=key, format="json", limit=PAGE_SIZE, offset=offset)
        url = BASE_URL + "?" + urllib.parse.urlencode(query)
        try:
            with urllib.request.urlopen(url, timeout=120) as resp:
                payload = json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as exc:
            sys.exit(f"AppLovin API returned {exc.code}: {exc.read().decode('utf-8', 'replace')[:500]}")
        batch = payload.get("results", payload if isinstance(payload, list) else [])
        rows.extend(batch)
        if len(batch) < PAGE_SIZE:
            return rows
        offset += PAGE_SIZE


def campaign_params(identifier):
    is_id = len(identifier) == 32 and all(c in "0123456789abcdef" for c in identifier.lower())
    return {"filter_campaign_id" if is_id else "filter_campaign": identifier}


def list_campaigns(cfg, start, end):
    rows = fetch({"start": str(start), "end": str(end), "columns": "campaign,campaign_id,cost"}, cfg["report_key"])
    rows.sort(key=lambda r: -num(r.get("cost")))
    print(f"Campaigns with data {start} -> {end}:\n")
    for r in rows:
        print(f"  {num(r.get('cost')):>12,.2f}   {r.get('campaign')}   [{r.get('campaign_id')}]")


# ------------------------------------------------------------------ shaping


def num(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def clean_name(name):
    """Drop the app-name prefix ahead of the leading V_/P_ token, if there is one.

    Must anchor on a separator: concept codes like V_Concept 01_003P_Sam contain
    an internal 'P_', and a bare find() chops the name in half there.
    """
    if name.startswith(("V_", "P_")):
        return name
    match = re.search(r"[_ ]([VP]_)", name)
    return name[match.start(1):] if match else name


def creative_type(name):
    if name.startswith("V_") or name.lower().endswith((".mp4", ".mov")):
        return "Video"
    if name.startswith("P_") or name.lower().endswith(".html"):
        return "Playable"
    return "Image"


def parse_date(value):
    value = (value or "").strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%b-%Y"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def first_serve_dates(identifiers, key, end, wanted):
    """name -> earliest day the asset served in these campaigns.

    The Total sheet is the authority on Added Date; this only fills gaps.
    It is a first-SERVE date, not an upload date - an asset uploaded on the
    Monday but first delivered on the Thursday comes back as Thursday.

    assetAnalyticsReport only serves the last 45 days, so a creative whose
    first impression predates that window cannot be recovered here and is
    left blank for manual entry.
    """
    if not wanted:
        return {}
    start = max(date.today() - timedelta(days=LOOKBACK_DAYS), end - timedelta(days=LOOKBACK_DAYS))
    found = {}
    for identifier in identifiers:
        rows = fetch(
            dict(
                {"start": str(start), "end": str(end),
                 "columns": "day,asset_name,impressions"},
                **campaign_params(identifier),
            ),
            key,
        )
        for row in rows:
            if not num(row.get("impressions")):
                continue
            name = clean_name(str(row.get("asset_name", "")).strip())
            if name not in wanted:
                continue
            when = parse_date(str(row.get("day", "")))
            if when and (name not in found or when < found[name]):
                found[name] = when
    return found


# --------------------------------------------------------------- Total sheet


def find_totals_file(totals_dir, label, other_labels=()):
    """Total sheet CSV for this label.

    A short label is a prefix of a longer one - BLD matches a file named
    ..._BLD_IOS_... as readily as ..._BLD_... - so any file that also
    matches a longer configured label belongs to that label, not this one.
    Without this, BLD silently reads BLD_IOS's dates.
    """
    totals_dir = os.path.expanduser(totals_dir)
    exact = os.path.join(totals_dir, f"{label}.csv")
    if os.path.exists(exact):
        return exact

    def hit(path, lbl):
        return re.search(rf"[_\-. ]{re.escape(lbl)}[_\-. ]",
                         os.path.basename(path), re.I)

    rivals = [l for l in other_labels
              if l.lower() != label.lower() and len(l) > len(label)]
    matches = []
    for p in glob.glob(os.path.join(totals_dir, "*.csv")):
        if not hit(p, label):
            continue
        if any(hit(p, r) for r in rivals):
            continue                      # belongs to the more specific label
        matches.append(p)

    matches.sort()
    if len(matches) > 1:
        print(f"  ! several Total files match {label}: "
              f"{', '.join(os.path.basename(m) for m in matches)}", file=sys.stderr)
        print(f"    using {os.path.basename(matches[0])} - rename the others "
              f"or split the label in config.json", file=sys.stderr)
    return matches[0] if matches else None


def read_totals(path):
    """name -> Added Date, plus the Total sheet's own row order."""
    added, order = {}, []
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.reader(fh)
        header = [h.strip() for h in next(reader)]
        try:
            ci, di = header.index("Creatives"), header.index("Added Date")
        except ValueError:
            print(f"  ! {os.path.basename(path)} has no Creatives/Added Date columns - skipped.", file=sys.stderr)
            return added, order
        for row in reader:
            if len(row) <= max(ci, di) or not row[ci].strip():
                continue
            name = clean_name(row[ci].strip())
            when = parse_date(row[di])
            if name not in added:
                added[name] = when
                order.append((creative_type(name), name))
    return added, order


def read_prior_order(path):
    """Row order from the newest weekly tab of a master workbook, if given."""
    order = []
    if not path:
        return order
    path = os.path.expanduser(path)
    if not os.path.exists(path):
        print(f"  ! workbook not found: {path} - falling back to Total sheet order.", file=sys.stderr)
        return order
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    tabs = [s for s in wb.sheetnames if TAB_RE.match(s)]
    if tabs:
        ws = wb[tabs[0]]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if "Creatives" in header:
            ci = header.index("Creatives")
            for row in ws.iter_rows(min_row=2):
                if row[ci].value:
                    name = clean_name(str(row[ci].value).strip())
                    order.append((creative_type(name), name))
    wb.close()
    return order


def order_rows(records, prior):
    """Keep the reference order; append newcomers to the end of their type block."""
    by_name = {r["name"]: r for r in records}
    ordered, used = [], set()
    for _, name in prior:
        if name in by_name and name not in used:
            ordered.append(by_name[name])
            used.add(name)
    for kind in TYPE_ORDER:
        newcomers = sorted(
            (r for r in records if r["type"] == kind and r["name"] not in used),
            key=lambda r: -r["cost"],
        )
        if not newcomers:
            continue
        last = max((i for i, r in enumerate(ordered) if r["type"] == kind), default=None)
        at = last + 1 if last is not None else len(ordered)
        ordered[at:at] = newcomers
        used.update(r["name"] for r in newcomers)
    if not prior:
        ordered.sort(key=lambda r: (TYPE_ORDER.index(r["type"]) if r["type"] in TYPE_ORDER else 9, -r["cost"]))
    return ordered


# --------------------------------------------------------------- workbook out


def write_tab(records, label, out_path, highlight):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = label
    ws.append(HEADER)
    for i, cell in enumerate(ws[1], start=1):
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS[i - 1]

    for r, rec in enumerate(records, start=2):
        ws.cell(r, 1, rec["type"])
        b = ws.cell(r, 2, rec["added"])
        b.number_format = "yyyy\\-mm\\-dd"
        if rec["added"] is None:
            b.fill = AMBER
        elif rec.get("derived"):
            b.fill = BLUE
        ws.cell(r, 3, rec["name"])
        ws.cell(r, 4, rec["cost"]).number_format = "\\$#,##0.00"
        ws.cell(r, 5, f"=IFERROR(G{r}/F{r},0)").number_format = "0.00%"
        ws.cell(r, 6, rec["impressions"])
        ws.cell(r, 7, rec["clicks"])
        ws.cell(r, 8, f"=IFERROR(D{r} / SUMIF($A:$A, A{r}, $D:$D),0)").number_format = "0%"
        for c in range(1, 9):
            ws.cell(r, c).font = BODY_FONT

    ws.freeze_panes = "A2"
    ws.conditional_formatting.add(
        f"D2:D{len(records) + 1}",
        CellIsRule(operator="greaterThan", formula=[str(highlight)], fill=PatternFill("solid", bgColor=GREEN)),
    )
    wb.save(out_path)


def write_csvs(records, outdir):
    for fname, kind in {"videos": "Video", "htmls": "Playable", "images": "Image"}.items():
        with open(os.path.join(outdir, f"{fname}.csv"), "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(["Name", "Status", "Upload time (UTC)", "Impressions", "Clicks", "CTR", "Spend"])
            for rec in records:
                if rec["type"] != kind:
                    continue
                ctr = rec["clicks"] / rec["impressions"] * 100 if rec["impressions"] else 0
                w.writerow([rec["name"], "", "", rec["impressions"], rec["clicks"], round(ctr, 2), round(rec["cost"], 2)])


# -------------------------------------------------------------------- main


def build(label, spec, cfg, start, end, outdir, want_csv, totals_dir,
          all_labels=(), no_lookup=False):
    identifiers = campaign_ids(spec, label)
    rows, per_id = [], []
    for identifier in identifiers:
        got = fetch(
            dict(
                {"start": str(start), "end": str(end), "columns": "asset_id,asset_name,impressions,clicks,cost"},
                **campaign_params(identifier),
            ),
            cfg["report_key"],
        )
        per_id.append((identifier, len(got)))
        rows.extend(got)
    if len(identifiers) > 1:
        print(f"\n{label}: merging {len(identifiers)} campaigns - " +
              ", ".join(f"{i} ({n} rows)" for i, n in per_id))
    if not rows:
        print(f"\n{label}: no rows for {', '.join(identifiers)} "
              f"between {start} and {end} - skipped.")
        return

    totals_path = find_totals_file(totals_dir, label, all_labels)
    if totals_path:
        added, totals_order = read_totals(totals_path)
        age = (date.today() - date.fromtimestamp(os.path.getmtime(totals_path))).days
    else:
        added, totals_order, age = {}, [], None

    merged = {}
    for row in rows:
        name = clean_name(str(row.get("asset_name", "")).strip())
        if not name:
            continue
        rec = merged.setdefault(name, {
            "name": name,
            "type": creative_type(name),
            "added": added.get(name),
            "cost": 0.0, "impressions": 0, "clicks": 0,
        })
        # one asset can run in several merged campaigns - sum, never replace
        rec["cost"] += num(row.get("cost"))
        rec["impressions"] += int(num(row.get("impressions")))
        rec["clicks"] += int(num(row.get("clicks")))
    records = list(merged.values())
    gaps = {r["name"] for r in records if r["added"] is None}
    derived = {}
    if gaps and not no_lookup:
        derived = first_serve_dates(identifiers, cfg["report_key"], end, gaps)
        for r in records:
            if r["added"] is None and r["name"] in derived:
                r["added"] = derived[r["name"]]
                r["derived"] = True

    records = order_rows(records, read_prior_order(spec.get("workbook")) or totals_order)

    week = week_label(start, end)
    os.makedirs(outdir, exist_ok=True)
    out_path = os.path.join(outdir, f"{label}_Applovin_Creative_Performance_{week}.xlsx")
    write_tab(records, week, out_path, spec.get("highlight", 1000))
    if want_csv:
        write_csvs(records, outdir)

    print(f"\n{label}  {week}  ->  {out_path}")
    for kind in TYPE_ORDER:
        block = [r for r in records if r["type"] == kind]
        if not block:
            continue
        spend = sum(r["cost"] for r in block)
        imps = sum(r["impressions"] for r in block)
        clicks = sum(r["clicks"] for r in block)
        print(
            f"  {kind:<9}{len(block):>4} creatives   ${spend:>11,.2f}"
            f"   CTR {(clicks / imps * 100 if imps else 0):>5.2f}%"
        )
    print(f"  {'total':<9}{len(records):>4} creatives   ${sum(r['cost'] for r in records):>11,.2f}")

    if not totals_path:
        print(f"  ! no Total sheet for {label} in {totals_dir} - every Added Date is blank.")
    elif age and age > 10:
        print(f"  ! {os.path.basename(totals_path)} was last changed {age} days ago - re-export it.")
    filled = [r for r in records if r.get("derived")]
    if filled:
        print(f"  {len(filled)} Added Date(s) filled from first serve "
              f"(blue cells) - first delivery, not upload:")
        for r in sorted(filled, key=lambda r: r["added"]):
            print(f"      {r['added']}  {r['name']}")

    missing = [r["name"] for r in records if r["added"] is None]
    if missing:
        print(f"  ! {len(missing)} creative(s) with no Added Date (amber) - "
              f"not in the Total sheet and no serve in the last {LOOKBACK_DAYS} days:")
        for name in missing:
            print(f"      {name}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--campaign", help="Config label (default: every configured campaign)")
    ap.add_argument("--week-end", help="Last day of the reporting week (YYYY-MM-DD). Default: last Tuesday")
    ap.add_argument("--outdir", default="weeks", help="Output folder (default: ./weeks/<week>)")
    ap.add_argument("--totals", default="totals", help="Folder holding the Total sheet CSVs (default: ./totals)")
    ap.add_argument("--config", default="config.json")
    ap.add_argument("--csv", action="store_true", help="Also write raw htmls/videos/images.csv")
    ap.add_argument("--no-date-lookup", action="store_true",
                    help="Skip the first-serve lookup for missing Added Dates")
    ap.add_argument("--list-campaigns", action="store_true")
    args = ap.parse_args()

    cfg = load_config(args.config)
    end = datetime.strptime(args.week_end, "%Y-%m-%d").date() if args.week_end else last_tuesday()
    start = end - timedelta(days=6)

    if (date.today() - start).days > 45:
        print(f"! {start} is outside the API's 45-day window - expect nothing back.", file=sys.stderr)

    if args.list_campaigns:
        list_campaigns(cfg, start, end)
        return

    labels = [args.campaign] if args.campaign else list(cfg["campaigns"])
    if not labels:
        sys.exit("No campaigns configured. Run --list-campaigns and fill in config.json.")

    outdir = os.path.join(args.outdir, week_label(start, end))
    for label in labels:
        build(
            label,
            cfg["campaigns"].get(label, {"id": label}),
            cfg,
            start,
            end,
            os.path.join(outdir, label) if args.csv else outdir,
            args.csv,
            args.totals,
            list(cfg["campaigns"]),
            args.no_date_lookup,
        )


if __name__ == "__main__":
    main()
